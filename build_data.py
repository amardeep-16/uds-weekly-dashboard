#!/usr/bin/env python3
"""build_data.py — parse UDS_PPT_Data_Wk_*.xlsx → data.js"""

import openpyxl, json, re, math, os, glob
from datetime import datetime

# ─── WINDOW config (front-anchored week-group count per sheet) ─────────────
WINDOW = {
    "cci": 13, "ei": 13, "source": 10, "client": 8, "pod": 13,
    "issue": 13, "zone": 8, "c2": 9, "postfake": 13,
}

SKIP_LABELS = {
    "Chatbot", "shadowfax_app", "Auto-Email", "whatsapp",
    "Express", "B2B", "Shadowfax", "Prime Large",
    "whatsapp_customer_response", "(blank)", "Fake Source UDS",
}

# ─── Core helpers ──────────────────────────────────────────────────────────
def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        if math.isnan(v) or math.isinf(v):
            return None
        return float(v)
    s = str(v).strip()
    if not s or s.startswith("#"):
        return None
    s = s.replace(",", "").replace("%", "")
    try:
        return float(s)
    except Exception:
        return None


def wk(label):
    if label is None:
        return None
    m = re.search(r"(\d+)", str(label).strip())
    if not m:
        return None
    return f"W{int(m.group(1)):02d}"


def clean_nan(obj):
    if isinstance(obj, float):
        return None if (math.isnan(obj) or math.isinf(obj)) else obj
    if isinstance(obj, dict):
        return {k: clean_nan(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [clean_nan(v) for v in obj]
    return obj


def nonempty_rows(rows):
    result = {}
    for label, wd in rows.items():
        if label in SKIP_LABELS:
            continue
        if any(v is not None for weekdata in wd.values() for v in weekdata.values()):
            result[label] = wd
    return result


def get_rows(ws, max_r, max_col=None):
    kw = {"min_row": 1, "max_row": max_r + 2, "values_only": True}
    if max_col:
        kw["max_col"] = max_col
    return list(ws.iter_rows(**kw))


def pod_c0(ws):
    """2nd column (0-indexed) whose row-index-1 value == 'POD'."""
    row = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
    pods = [i for i, v in enumerate(row) if v == "POD"]
    return pods[1] if len(pods) >= 2 else (pods[0] if pods else 0)


# ─── Block readers ─────────────────────────────────────────────────────────
def read_block(ws, r_top, r_sub, r0, r1, c0=0, maxgroups=13,
               ei_normalize=False, max_col=None):
    """
    Two-row merged-header block reader.
    Returns (rows_dict, weeks_list, all_sub_keys).
    rows_dict: {label: {wk_label: {sub_key: float|None}}}
    """
    max_r = max(r_top, r_sub, r1)
    all_rows = get_rows(ws, max_r, max_col)

    def _row(ri):
        return all_rows[ri] if ri < len(all_rows) else ()

    top_row = _row(r_top)
    sub_row = _row(r_sub)
    ncols = max(len(top_row), len(sub_row))

    # Detect group start cols in top_row (col > c0, non-empty)
    group_starts = [
        c for c in range(c0 + 1, ncols)
        if c < len(top_row) and top_row[c] is not None
        and str(top_row[c]).strip()
    ][:maxgroups]

    if not group_starts:
        return {}, [], []

    # Per-group sub-col detection (handles variable sub-count)
    groups = []
    for i, gc in enumerate(group_starts):
        next_gc = group_starts[i + 1] if i + 1 < len(group_starts) else ncols
        wk_label = wk(top_row[gc])
        raw = []
        for c in range(gc, min(next_gc, len(sub_row))):
            v = sub_row[c]
            if v is not None and str(v).strip():
                k = str(v).strip()
                if ei_normalize and k == "CCI":
                    k = "EI"
                raw.append((c, k))
        # De-duplicate sub-keys with #2, #3 …
        seen = {}
        deduped = []
        for c, k in raw:
            if k not in seen:
                seen[k] = 1
                deduped.append((c, k))
            else:
                seen[k] += 1
                deduped.append((c, f"{k}#{seen[k]}"))
        groups.append((wk_label, deduped))

    weeks = [g[0] for g in groups]

    # Read data rows
    rows = {}
    for ri in range(r0, r1 + 1):
        if ri >= len(all_rows):
            break
        dr = _row(ri)
        raw_label = dr[c0] if c0 < len(dr) else None
        if raw_label is None:
            continue
        label = str(raw_label).strip()
        if not label:
            continue
        rows[label] = {}
        for wk_label, sub_keys in groups:
            rows[label][wk_label] = {}
            for c, k in sub_keys:
                v = dr[c] if c < len(dr) else None
                rows[label][wk_label][k] = num(v)

    all_subs = []
    for _, sub_keys in groups:
        for _, k in sub_keys:
            if k not in all_subs:
                all_subs.append(k)

    return rows, weeks, all_subs


def read_simple(ws, r_hdr, r0, r1, c0=0, ncols=None):
    """Single-header-row reader. Returns (rows_dict, col_headers)."""
    max_r = max(r_hdr, r1)
    all_rows = get_rows(ws, max_r)
    hdr = all_rows[r_hdr] if r_hdr < len(all_rows) else ()

    col_headers, col_indices = [], []
    for c in range(c0 + 1, len(hdr)):
        if hdr[c] is not None and str(hdr[c]).strip():
            col_headers.append(str(hdr[c]).strip())
            col_indices.append(c)

    if ncols:
        col_headers = col_headers[:ncols]
        col_indices = col_indices[:ncols]

    rows = {}
    for ri in range(r0, r1 + 1):
        if ri >= len(all_rows):
            break
        dr = all_rows[ri]
        raw_label = dr[c0] if c0 < len(dr) else None
        if raw_label is None:
            continue
        label = str(raw_label).strip()
        if not label:
            continue
        rows[label] = {}
        for c, h in zip(col_indices, col_headers):
            v = dr[c] if c < len(dr) else None
            rows[label][h] = num(v)

    return rows, col_headers


# ─── KPI & chart builder helpers ───────────────────────────────────────────
def _kpi(label, value, fmt, delta=None, deltafmt=None, goodlow=True, suffix=None, text=None):
    k = {"label": label, "value": value, "fmt": fmt, "goodlow": goodlow}
    if delta is not None:
        k["delta"] = delta
        k["deltafmt"] = deltafmt or "pp"
    if suffix:
        k["suffix"] = suffix
    if text:
        k["text"] = text
    return k


def _get(rows, label, week, metric):
    return rows.get(label, {}).get(week, {}).get(metric)


def _series(rows, label, weeks, metric):
    return [_get(rows, label, w, metric) for w in weeks]


def _delta(rows, label, metric, weeks):
    vals = []
    for w in reversed(weeks):
        v = _get(rows, label, w, metric)
        if v is not None:
            vals.append(v)
        if len(vals) == 2:
            break
    if len(vals) < 2:
        return None
    return vals[0] - vals[1]


def _latest_val(rows, label, weeks, metric):
    for w in reversed(weeks):
        v = _get(rows, label, w, metric)
        if v is not None:
            return v
    return None


def _line(title, fmt, weeks, datasets):
    return {
        "type": "line", "title": title, "fmt": fmt,
        "labels": weeks,
        "datasets": [{"label": lb, "data": data} for lb, data in datasets],
    }


def _bar(title, fmt, labels, datasets, horizontal=False):
    return {
        "type": "bar", "title": title, "fmt": fmt, "horizontal": horizontal,
        "labels": labels,
        "datasets": [{"label": lb, "data": data} for lb, data in datasets],
    }


def _doughnut(title, fmt, labels, data):
    return {
        "type": "doughnut", "title": title, "fmt": fmt,
        "labels": labels,
        "datasets": [{"label": "", "data": data}],
    }


def _table(title, headers, rows, fmt):
    return {"title": title, "headers": headers, "rows": rows, "fmt": fmt}


# ─── Sheet 1: CCI Slide ────────────────────────────────────────────────────
def parse_cci(wb):
    ws = wb["CCI Slide"]
    W = WINDOW["cci"]
    # Overall-All: top=20, sub=21, data=22-24, c0=0
    rows_all, weeks, _ = read_block(ws, 20, 21, 22, 24, c0=0, maxgroups=W)
    # Non-Meesho: top=27, sub=28, data=29-31
    rows_nm, _, _ = read_block(ws, 27, 28, 29, 31, c0=0, maxgroups=W)
    # Meesho: top=34, sub=35, data=36-38
    rows_m, _, _ = read_block(ws, 34, 35, 36, 38, c0=0, maxgroups=W)
    # Complaint sources Overall — simple header=41, data=42-46
    rows_src, src_weeks = read_simple(ws, 41, 42, 46, c0=0, ncols=W)

    lw = weeks[-1] if weeks else "?"
    prev_w = weeks[-2] if len(weeks) > 1 else None

    def cval(r, label, metric):
        return _latest_val(r, label, weeks, metric)

    def cdelta(r, label, metric):
        return _delta(r, label, metric, weeks)

    ov_fwd = cval(rows_all, "FWD", "CCI") or cval(rows_all, "Fwd", "CCI")
    ov_rev = cval(rows_all, "REV", "CCI") or cval(rows_all, "Rev", "CCI")
    nm_ov  = cval(rows_nm, "Overall", "CCI")
    m_ov   = cval(rows_m, "Overall", "CCI")

    d_fwd = cdelta(rows_all, "FWD", "CCI") or cdelta(rows_all, "Fwd", "CCI")
    d_rev = cdelta(rows_all, "REV", "CCI") or cdelta(rows_all, "Rev", "CCI")
    d_nm  = cdelta(rows_nm, "Overall", "CCI")
    d_m   = cdelta(rows_m, "Overall", "CCI")

    kpis = [
        _kpi(f"CCI Fwd {lw}", ov_fwd, "pct2", d_fwd, "pp", goodlow=True),
        _kpi(f"CCI Rev {lw}", ov_rev, "pct2", d_rev, "pp", goodlow=True),
        _kpi(f"Non-Meesho CCI {lw}", nm_ov, "pct2", d_nm, "pp", goodlow=True),
        _kpi(f"Meesho CCI {lw}", m_ov, "pct2", d_m, "pp", goodlow=True),
    ]

    # Charts
    # 1. CCI trend (Overall section: Fwd/Rev/Overall)
    chart_labels_map = {"FWD": "Fwd", "REV": "Rev", "Overall": "Overall",
                        "Fwd": "Fwd", "Rev": "Rev"}
    trend_ds = []
    for raw_lbl in ("FWD", "Fwd", "REV", "Rev", "Overall"):
        if raw_lbl in rows_all:
            trend_ds.append((chart_labels_map.get(raw_lbl, raw_lbl),
                             _series(rows_all, raw_lbl, weeks, "CCI")))
    seen = set()
    trend_ds = [d for d in trend_ds if d[0] not in seen and not seen.add(d[0])]

    # 2. CCI latest-week comparison: All / NM / Meesho for Fwd+Rev
    bar_labels = ["All Fwd", "All Rev", "NM Fwd", "NM Rev", "Meesho Fwd", "Meesho Rev"]
    def _lv(r, l, m):
        v = _latest_val(r, l, weeks, m)
        return v
    bar_data = [
        _lv(rows_all, "FWD", "CCI") or _lv(rows_all, "Fwd", "CCI"),
        _lv(rows_all, "REV", "CCI") or _lv(rows_all, "Rev", "CCI"),
        _lv(rows_nm, "FWD", "CCI") or _lv(rows_nm, "Fwd", "CCI"),
        _lv(rows_nm, "REV", "CCI") or _lv(rows_nm, "Rev", "CCI"),
        _lv(rows_m, "FWD", "CCI") or _lv(rows_m, "Fwd", "CCI"),
        _lv(rows_m, "REV", "CCI") or _lv(rows_m, "Rev", "CCI"),
    ]

    # 3. Complaint source doughnut (latest available week)
    src_lw = src_weeks[-1] if src_weeks else None
    src_mix_labels = [l for l in rows_src if l != "Grand Total" and l not in SKIP_LABELS]
    src_mix_data = [rows_src.get(l, {}).get(src_lw) for l in src_mix_labels]

    charts = [
        _line("CCI Trend – Overall", "pct2", weeks, trend_ds),
        _bar("CCI Latest Week Comparison", "pct2", bar_labels, [("CCI %", bar_data)]),
        _doughnut(f"Complaint Source Mix {src_lw}", "int", src_mix_labels, src_mix_data),
    ]

    # Table: CCI All for last 5 weeks
    tw = weeks[-5:] if len(weeks) >= 5 else weeks
    tbl_rows = []
    for lbl in ("FWD", "Fwd", "REV", "Rev", "Overall"):
        if lbl in rows_all:
            tbl_rows.append([lbl] + [_get(rows_all, lbl, w, "CCI") for w in tw])
    if tbl_rows:
        tables = [_table("CCI % – All POD", ["Type"] + tw, tbl_rows,
                         ["t"] + ["pct2"] * len(tw))]
    else:
        tables = []

    ov_v = _latest_val(rows_all, "Overall", weeks, "CCI")
    ov_d = _delta(rows_all, "Overall", "CCI", weeks)
    ov_pct = f"{(ov_v or 0)*100:.2f}%"
    ov_pp  = f"{(ov_d or 0)*100:+.2f}pp" if ov_d else ""

    return {
        "id": "cci", "name": "CCI Slide", "icon": "CI",
        "subtitle": "Complaint Contact Index by week & segment",
        "latest_week": lw,
        "narrative": (f"{lw}: Overall CCI at {ov_pct} ({ov_pp} vs prior week). "
                      f"Non-Meesho at {(nm_ov or 0)*100:.2f}%, "
                      f"Meesho at {(m_ov or 0)*100:.2f}%."),
        "kpis": kpis, "charts": charts, "tables": tables,
    }


# ─── Sheet 2: NEW EI% ──────────────────────────────────────────────────────
def parse_ei(wb):
    ws = wb["NEW EI%"]
    W = WINDOW["ei"]
    rows_all, weeks, _ = read_block(ws, 0, 1, 2, 4, c0=0, maxgroups=W, ei_normalize=True)
    rows_nm,  _, _     = read_block(ws, 7, 8, 9, 11, c0=0, maxgroups=W, ei_normalize=True)
    rows_m,   _, _     = read_block(ws, 14, 15, 16, 18, c0=0, maxgroups=W, ei_normalize=True)
    rows_src, src_wks  = read_simple(ws, 35, 36, 44, c0=0, ncols=W)

    lw = weeks[-1] if weeks else "?"

    def ev(r, l): return _latest_val(r, l, weeks, "EI")
    def ed(r, l): return _delta(r, l, "EI", weeks)

    kpis = [
        _kpi(f"EI Fwd {lw}",  ev(rows_all,"Fwd"), "num2", ed(rows_all,"Fwd"), "num2",
             goodlow=True, suffix="/1k"),
        _kpi(f"EI Rev {lw}",  ev(rows_all,"Rev"), "num2", ed(rows_all,"Rev"), "num2",
             goodlow=True, suffix="/1k"),
        _kpi(f"Non-Meesho EI {lw}", ev(rows_nm,"Overall"), "num2", ed(rows_nm,"Overall"),
             "num2", goodlow=True, suffix="/1k"),
        _kpi(f"Meesho EI {lw}", ev(rows_m,"Overall"), "num2", ed(rows_m,"Overall"),
             "num2", goodlow=True, suffix="/1k"),
    ]

    trend_ds = [(l, _series(rows_all, l, weeks, "EI")) for l in ("Fwd","Rev","Overall") if l in rows_all]

    seg_labels = ["All Fwd","All Rev","All Ovr","NM Fwd","NM Rev","NM Ovr","Msho Fwd","Msho Rev","Msho Ovr"]
    seg_vals = [
        ev(rows_all,"Fwd"), ev(rows_all,"Rev"), ev(rows_all,"Overall"),
        ev(rows_nm,"Fwd"),  ev(rows_nm,"Rev"),  ev(rows_nm,"Overall"),
        ev(rows_m,"Fwd"),   ev(rows_m,"Rev"),   ev(rows_m,"Overall"),
    ]

    src_lw = src_wks[-1] if src_wks else None
    src_labs = [l for l in rows_src if l not in ("Grand Total",) and l not in SKIP_LABELS]
    src_data = [rows_src.get(l,{}).get(src_lw) for l in src_labs]

    charts = [
        _line("EI Trend – Overall", "num2", weeks, trend_ds),
        _bar("EI by Segment (Latest Week)", "num2", seg_labels, [("EI /1k", seg_vals)]),
        _doughnut(f"Complaint Source Mix {src_lw}", "int", src_labs, src_data),
    ]

    tw = weeks[-5:] if len(weeks) >= 5 else weeks
    tbl_rows = [[l] + [_get(rows_all,l,w,"EI") for w in tw]
                for l in ("Fwd","Rev","Overall") if l in rows_all]
    tables = [_table("EI – Overall", ["Type"]+tw, tbl_rows, ["t"]+["num2"]*len(tw))]

    ov = ev(rows_all,"Overall")
    od = ed(rows_all,"Overall")
    return {
        "id":"ei","name":"NEW EI%","icon":"EI",
        "subtitle":"Escalation Index per 1 000 shipments",
        "latest_week": lw,
        "narrative": (f"{lw}: Overall EI = {ov:.2f}/1k ({od:+.2f} vs prior)." if ov and od
                      else f"{lw}: Overall EI = {ov:.2f}/1k." if ov else f"{lw}: EI data."),
        "kpis":kpis,"charts":charts,"tables":tables,
    }


# ─── Sheet 3: Source Wise Del & Pik ────────────────────────────────────────
def parse_source(wb):
    ws = wb["Source Wise Del & Pik"]
    W = WINDOW["source"]
    rows_fwd, weeks_f, _ = read_block(ws, 1, 2, 3, 11, c0=0, maxgroups=W)
    rows_rev, weeks_r, _ = read_block(ws, 13, 14, 15, 21, c0=0, maxgroups=W)

    rows_fwd = nonempty_rows(rows_fwd)
    rows_rev = nonempty_rows(rows_rev)
    lw  = weeks_f[-1] if weeks_f else "?"
    lwr = weeks_r[-1] if weeks_r else "?"

    def v(r, l, w, m): return _get(r, l, w, m)

    fwd_uds_gt = _latest_val(rows_fwd,"Grand Total",weeks_f,"UDS")
    fwd_dl_gt  = _latest_val(rows_fwd,"Grand Total",weeks_f,"DL")
    rev_uds_gt = _latest_val(rows_rev,"Grand Total",weeks_r,"UDS")
    rev_pk_gt  = _latest_val(rows_rev,"Grand Total",weeks_r,"Picked") or \
                 _latest_val(rows_rev,"Grand Total",weeks_r,"Picked / QC")

    kpis = [
        _kpi(f"Fwd UDS {lw}",    fwd_uds_gt, "int", goodlow=False),
        _kpi(f"Fwd DL {lw}",     fwd_dl_gt,  "int", goodlow=False),
        _kpi(f"Rev UDS {lw}",    rev_uds_gt, "int", goodlow=False),
        _kpi(f"Rev Picked {lwr}", rev_pk_gt,  "int", goodlow=False),
    ]

    # Trend chart: Forward UDS by source
    src_labels_fwd = [l for l in rows_fwd if l != "Grand Total"]
    trend_ds_fwd = [(l, _series(rows_fwd, l, weeks_f, "UDS")) for l in src_labels_fwd]

    # Bar: latest-week forward breakdown by source
    bar_labs = src_labels_fwd
    bar_uds  = [_latest_val(rows_fwd, l, weeks_f, "UDS") for l in bar_labs]
    bar_dl   = [_latest_val(rows_fwd, l, weeks_f, "DL")  for l in bar_labs]

    # Doughnut: rev latest-week UDS mix
    rev_labs = [l for l in rows_rev if l != "Grand Total"]
    rev_data = [_latest_val(rows_rev, l, weeks_r, "UDS") for l in rev_labs]

    charts = [
        _line("Fwd UDS Trend by Source", "int", weeks_f, trend_ds_fwd),
        _bar(f"Fwd Latest Week {lw} – UDS vs DL", "int", bar_labs,
             [("UDS", bar_uds), ("DL", bar_dl)]),
        _doughnut(f"Rev UDS Mix {lwr}", "int", rev_labs, rev_data),
    ]

    tw = weeks_f[-5:] if len(weeks_f) >= 5 else weeks_f
    tbl_rows = []
    for l in src_labels_fwd + ["Grand Total"]:
        if l in rows_fwd:
            row = [l]
            for w in tw:
                row += [_get(rows_fwd,l,w,"UDS"), _get(rows_fwd,l,w,"DL"),
                        _get(rows_fwd,l,w,"RTO")]
            tbl_rows.append(row)
    hdrs = ["Source"] + [f"{w} UDS / DL / RTO" for w in tw]
    tbl_flat_hdrs = ["Source"]
    tbl_fmt = ["t"]
    for w in tw:
        tbl_flat_hdrs += [f"{w}·UDS", f"{w}·DL", f"{w}·RTO"]
        tbl_fmt += ["int","int","int"]
    flat_rows = []
    for l in src_labels_fwd + ["Grand Total"]:
        if l in rows_fwd:
            r = [l]
            for w in tw:
                r += [_get(rows_fwd,l,w,"UDS"),_get(rows_fwd,l,w,"DL"),_get(rows_fwd,l,w,"RTO")]
            flat_rows.append(r)
    tables = [_table("Forward Source UDS / DL / RTO", tbl_flat_hdrs, flat_rows, tbl_fmt)]

    return {
        "id":"source","name":"Source Wise Del & Pik","icon":"SW",
        "subtitle":"UDS delivery & pickup by complaint source",
        "latest_week": lw,
        "narrative": (f"{lw}: Fwd Grand Total — UDS {int(fwd_uds_gt or 0):,}, "
                      f"DL {int(fwd_dl_gt or 0):,}."),
        "kpis":kpis,"charts":charts,"tables":tables,
    }


# ─── Sheet 4: Client Wise Del & Pik ───────────────────────────────────────
def parse_client(wb):
    ws = wb["Client Wise Del & Pik"]
    W = WINDOW["client"]
    rows_fwd, weeks, subs_f = read_block(ws, 0, 1, 2, 10, c0=0, maxgroups=W)
    rows_rev, weeks_r, subs_r = read_block(ws, 12, 13, 14, 17, c0=0, maxgroups=W)

    rows_fwd = nonempty_rows(rows_fwd)
    rows_rev = nonempty_rows(rows_rev)
    lw = weeks[-1] if weeks else "?"

    # % = Delivered%, %#2 = RTO%
    del_pct_k = "%" if "%" in subs_f else subs_f[2] if len(subs_f) > 2 else "%"
    rto_pct_k = "%#2" if "%#2" in subs_f else subs_f[4] if len(subs_f) > 4 else "%#2"

    gt_del = _latest_val(rows_fwd,"Grand Total",weeks,del_pct_k)
    gt_rto = _latest_val(rows_fwd,"Grand Total",weeks,rto_pct_k)
    gt_del_d = _delta(rows_fwd,"Grand Total",del_pct_k,weeks)
    gt_rto_d = _delta(rows_fwd,"Grand Total",rto_pct_k,weeks)

    rev_lw = weeks_r[-1] if weeks_r else "?"
    # Rev sub-keys: Picked%, CAN/QC%, etc
    rev_pick_k = subs_r[1] if len(subs_r) > 1 else "Picked"
    rev_pick_pk = subs_r[2] if len(subs_r) > 2 else "%"
    gt_pick = _latest_val(rows_rev,"Grand Total",weeks_r,rev_pick_pk)
    gt_pick_d = _delta(rows_rev,"Grand Total",rev_pick_pk,weeks_r)

    kpis = [
        _kpi(f"Fwd Del% {lw}",  gt_del, "pct2", gt_del_d, "pp", goodlow=False),
        _kpi(f"Fwd RTO% {lw}",  gt_rto, "pct2", gt_rto_d, "pp", goodlow=True),
        _kpi(f"Rev Pick% {rev_lw}", gt_pick, "pct2", gt_pick_d, "pp", goodlow=False),
        _kpi(f"Fwd Clients {lw}", len([l for l in rows_fwd if l!="Grand Total"]),
             "int", goodlow=False),
    ]

    client_labs = [l for l in rows_fwd if l != "Grand Total"]
    del_vals = [_latest_val(rows_fwd,l,weeks,del_pct_k) for l in client_labs]
    rto_vals = [_latest_val(rows_fwd,l,weeks,rto_pct_k) for l in client_labs]

    # Trend: Grand Total delivered%
    trend_ds = [("Del%", _series(rows_fwd,"Grand Total",weeks,del_pct_k)),
                ("RTO%", _series(rows_fwd,"Grand Total",weeks,rto_pct_k))]

    charts = [
        _bar(f"Client Del% {lw}", "pct2", client_labs, [("Del%", del_vals)]),
        _bar(f"Client RTO% {lw}", "pct2", client_labs, [("RTO%", rto_vals)]),
        _line("Fwd Del%/RTO% Trend – Grand Total", "pct2", weeks, trend_ds),
    ]

    tbl_rows = [[l, _latest_val(rows_fwd,l,weeks,"Grand Total"),
                 _latest_val(rows_fwd,l,weeks,del_pct_k),
                 _latest_val(rows_fwd,l,weeks,rto_pct_k)]
                for l in client_labs + ["Grand Total"] if l in rows_fwd]
    tables = [_table(f"Client Fwd Summary {lw}",
                     ["Client","Grand Total","Del%","RTO%"],
                     tbl_rows, ["t","int","pct2","pct2"])]

    return {
        "id":"client","name":"Client Wise Del & Pik","icon":"CL",
        "subtitle":"Delivery & RTO rates by client type",
        "latest_week": lw,
        "narrative":(f"{lw}: Overall Delivered {(gt_del or 0)*100:.1f}%, "
                     f"RTO {(gt_rto or 0)*100:.1f}%."),
        "kpis":kpis,"charts":charts,"tables":tables,
    }


# ─── Sheet 5: POD Wise CCI_Del, Pik & Fake % ──────────────────────────────
def parse_pod(wb):
    ws = wb["POD Wise CCI_Del, Pik & Fake %"]
    W = WINDOW["pod"]
    c0 = pod_c0(ws)
    max_col = c0 + W * 4 + 20
    rows, weeks, subs = read_block(ws, 1, 2, 3, 20, c0=c0, maxgroups=W, max_col=max_col)
    rows = nonempty_rows(rows)
    lw = weeks[-1] if weeks else "?"

    del_k  = "Delivered %" if "Delivered %" in subs else (subs[1] if len(subs)>1 else "Delivered %")
    fake_k = "Fake %" if "Fake %" in subs else (subs[2] if len(subs)>2 else "Fake %")
    uds_k  = "UDS" if "UDS" in subs else (subs[0] if subs else "UDS")

    gt_del  = _latest_val(rows,"Grand Total",weeks,del_k)
    gt_fake = _latest_val(rows,"Grand Total",weeks,fake_k)
    gt_uds  = _latest_val(rows,"Grand Total",weeks,uds_k)
    d_del   = _delta(rows,"Grand Total",del_k,weeks)
    d_fake  = _delta(rows,"Grand Total",fake_k,weeks)

    kpis = [
        _kpi(f"Del% GT {lw}",  gt_del,  "pct2", d_del,  "pp", goodlow=False),
        _kpi(f"Fake% GT {lw}", gt_fake, "pct2", d_fake, "pp", goodlow=True),
        _kpi(f"UDS GT {lw}",   gt_uds,  "pct2", goodlow=False),
        _kpi(f"PODs tracked",  len([l for l in rows if l!="Grand Total"]), "int", goodlow=False),
    ]

    pod_labs = [l for l in rows if l != "Grand Total"]
    del_vals  = [_latest_val(rows,l,weeks,del_k)  for l in pod_labs]
    fake_vals = [_latest_val(rows,l,weeks,fake_k) for l in pod_labs]

    trend_ds = [(del_k, _series(rows,"Grand Total",weeks,del_k))]
    if fake_k in subs:
        trend_ds.append((fake_k, _series(rows,"Grand Total",weeks,fake_k)))

    charts = [
        _bar(f"Delivered% by POD {lw}", "pct2", pod_labs, [("Del%", del_vals)]),
        _bar(f"Fake% by POD {lw}",      "pct2", pod_labs, [("Fake%", fake_vals)]),
        _line("Grand Total Trend – Del% & Fake%", "pct2", weeks, trend_ds),
    ]

    tbl_rows = [[l, _latest_val(rows,l,weeks,uds_k),
                 _latest_val(rows,l,weeks,del_k),
                 _latest_val(rows,l,weeks,fake_k)]
                for l in pod_labs + ["Grand Total"] if l in rows]
    tables = [_table(f"POD Summary {lw}",
                     ["POD","UDS","Del%","Fake%"],
                     tbl_rows, ["t","pct2","pct2","pct2"])]

    return {
        "id":"pod","name":"POD Wise CCI_Del, Pik & Fake %","icon":"PD",
        "subtitle":"Delivery, UDS and Fake % by Pincode-OD",
        "latest_week": lw,
        "narrative":(f"{lw}: Grand Total Del% = {(gt_del or 0)*100:.1f}%, "
                     f"Fake% = {(gt_fake or 0)*100:.1f}%."),
        "kpis":kpis,"charts":charts,"tables":tables,
    }


# ─── Sheet 6: Issue Type Wise_Fwd & Rev ───────────────────────────────────
def parse_issue(wb):
    ws = wb["Issue Type Wise_Fwd & Rev"]
    W = WINDOW["issue"]
    rows_fwd, weeks_f, subs_f = read_block(ws, 0, 1, 2, 10, c0=0, maxgroups=W)
    rows_rev, weeks_r, subs_r = read_block(ws, 12, 13, 14, 21, c0=0, maxgroups=W)

    rows_fwd = nonempty_rows(rows_fwd)
    rows_rev = nonempty_rows(rows_rev)
    lw  = weeks_f[-1] if weeks_f else "?"
    lwr = weeks_r[-1] if weeks_r else "?"

    awb_k = "AWB Count" if "AWB Count" in subs_f else subs_f[0] if subs_f else "AWB Count"

    gt_fwd = _latest_val(rows_fwd,"Grand Total",weeks_f,awb_k)
    gt_rev = _latest_val(rows_rev,"Grand Total",weeks_r,awb_k)

    issue_fwd = [(l, _latest_val(rows_fwd,l,weeks_f,awb_k))
                 for l in rows_fwd if l!="Grand Total"]
    issue_rev = [(l, _latest_val(rows_rev,l,weeks_r,awb_k))
                 for l in rows_rev if l!="Grand Total"]
    issue_fwd_s = sorted(issue_fwd, key=lambda x: x[1] or 0, reverse=True)[:10]
    issue_rev_s = sorted(issue_rev, key=lambda x: x[1] or 0, reverse=True)[:8]

    top_fwd = issue_fwd_s[0] if issue_fwd_s else ("N/A", None)
    top_rev = issue_rev_s[0] if issue_rev_s else ("N/A", None)

    kpis = [
        _kpi(f"Top Fwd Issue {lw}", top_fwd[1], "int", goodlow=True,
             text=top_fwd[0]),
        _kpi(f"Fwd Total AWB {lw}", gt_fwd, "int", goodlow=True),
        _kpi(f"Top Rev Issue {lwr}", top_rev[1], "int", goodlow=True,
             text=top_rev[0]),
        _kpi(f"Rev Total AWB {lwr}", gt_rev, "int", goodlow=True),
    ]

    fwd_labs = [x[0] for x in issue_fwd_s]
    fwd_data = [x[1] for x in issue_fwd_s]
    rev_labs = [x[0] for x in issue_rev_s]
    rev_data = [x[1] for x in issue_rev_s]

    charts = [
        _bar(f"Fwd Issue Types by AWB {lw}", "int", fwd_labs,
             [("AWB Count", fwd_data)], horizontal=True),
        _doughnut(f"Fwd Issue Mix {lw}", "int", fwd_labs, fwd_data),
        _bar(f"Rev Issue Types by AWB {lwr}", "int", rev_labs,
             [("AWB Count", rev_data)], horizontal=True),
    ]

    tbl_rows_f = [[l, _latest_val(rows_fwd,l,weeks_f,awb_k),
                   _latest_val(rows_fwd,l,weeks_f,"UDS  Contribution%") or
                   _latest_val(rows_fwd,l,weeks_f,"UDS Contribution%")]
                  for l in (fwd_labs + ["Grand Total"]) if l in rows_fwd]
    tables = [_table(f"Fwd Issue AWB {lw}", ["Issue Type","AWB Count","Contribution%"],
                     tbl_rows_f, ["t","int","pct2"])]

    return {
        "id":"issue","name":"Issue Type Wise_Fwd & Rev","icon":"IT",
        "subtitle":"AWB counts by complaint issue type",
        "latest_week": lw,
        "narrative":(f"{lw}: Fwd total AWB {int(gt_fwd or 0):,}; "
                     f"top issue: {top_fwd[0]}."),
        "kpis":kpis,"charts":charts,"tables":tables,
    }


# ─── Sheet 7: ZONE WISE ────────────────────────────────────────────────────
def parse_zone(wb):
    ws = wb["ZONE WISE"]
    W = WINDOW["zone"]
    rows_f, weeks_f, subs_f = read_block(ws, 0, 1, 2, 6, c0=0, maxgroups=W)
    rows_r, weeks_r, subs_r = read_block(ws, 9, 10, 11, 15, c0=0, maxgroups=W)

    lw  = weeks_f[-1] if weeks_f else "?"
    lwr = weeks_r[-1] if weeks_r else "?"

    del_k  = "DELIVERED"
    rto_k  = "RTO-RTS"
    fake_k = "Fake %"
    pick_k = "Picked"

    gt_del  = _latest_val(rows_f,"Grand Total",weeks_f,del_k)
    gt_rto  = _latest_val(rows_f,"Grand Total",weeks_f,rto_k)
    gt_fake = _latest_val(rows_f,"Grand Total",weeks_f,fake_k)
    gt_pick = _latest_val(rows_r,"Grand Total",weeks_r,pick_k)

    d_del   = _delta(rows_f,"Grand Total",del_k,weeks_f)
    d_fake  = _delta(rows_f,"Grand Total",fake_k,weeks_f)

    kpis = [
        _kpi(f"Fwd Del% {lw}",  gt_del,  "pct2", d_del,  "pp", goodlow=False),
        _kpi(f"Fwd RTO% {lw}",  gt_rto,  "pct2", goodlow=True),
        _kpi(f"Fwd Fake% {lw}", gt_fake, "pct2", d_fake, "pp", goodlow=True),
        _kpi(f"Rev Pick% {lwr}", gt_pick, "pct2", goodlow=False),
    ]

    zones = [l for l in rows_f if l != "Grand Total"]
    del_vals  = [_latest_val(rows_f,l,weeks_f,del_k)  for l in zones]
    fake_vals = [_latest_val(rows_f,l,weeks_f,fake_k) for l in zones]

    zones_r = [l for l in rows_r if l != "Grand Total"]
    pick_vals = [_latest_val(rows_r,l,weeks_r,pick_k) for l in zones_r]

    trend_ds = [("Del%", _series(rows_f,"Grand Total",weeks_f,del_k)),
                ("RTO%", _series(rows_f,"Grand Total",weeks_f,rto_k))]
    if fake_k in subs_f:
        trend_ds.append(("Fake%", _series(rows_f,"Grand Total",weeks_f,fake_k)))

    charts = [
        _bar(f"Fwd Del% by Zone {lw}",  "pct2", zones,   [("Del%",  del_vals)]),
        _bar(f"Rev Pick% by Zone {lwr}", "pct2", zones_r, [("Pick%", pick_vals)]),
        _line("Fwd Trend – Del%/RTO%/Fake%", "pct2", weeks_f, trend_ds),
    ]

    tbl_rows = [[z,
                 _latest_val(rows_f,z,weeks_f,del_k),
                 _latest_val(rows_f,z,weeks_f,rto_k),
                 _latest_val(rows_f,z,weeks_f,fake_k)]
                for z in zones + ["Grand Total"] if z in rows_f]
    tables = [_table(f"Zone Fwd {lw}",
                     ["Zone","Del%","RTO-RTS%","Fake%"],
                     tbl_rows, ["t","pct2","pct2","pct2"])]

    return {
        "id":"zone","name":"ZONE WISE","icon":"ZW",
        "subtitle":"Delivery & pickup performance by logistics zone",
        "latest_week": lw,
        "narrative":(f"{lw}: Fwd Del% {(gt_del or 0)*100:.1f}%, "
                     f"Fake% {(gt_fake or 0)*100:.1f}%."),
        "kpis":kpis,"charts":charts,"tables":tables,
    }


# ─── Sheet 8: Zone Wise - C2 ───────────────────────────────────────────────
def parse_c2(wb):
    ws = wb["Zone Wise - C2"]
    W = WINDOW["c2"]
    rows_all, weeks_a, subs_a = read_block(ws, 1, 2, 3, 8, c0=0, maxgroups=W)
    rows_cli, weeks_c, subs_c = read_block(ws, 12, 13, 14, 19, c0=0, maxgroups=W)
    rows_ei,  ei_wks          = read_simple(ws, 23, 24, 25, c0=0, ncols=W)

    lw  = weeks_a[-1] if weeks_a else "?"
    ei_k = "EI"

    gt_ei_all = _latest_val(rows_all,"Grand Total",weeks_a,ei_k)
    gt_ei_cli = _latest_val(rows_cli,"Grand Total",weeks_c,ei_k)
    d_ei_all  = _delta(rows_all,"Grand Total",ei_k,weeks_a)
    gt_esc    = _latest_val(rows_all,"Grand Total",weeks_a,"Total Esc")

    kpis = [
        _kpi(f"C2 EI All-Src {lw}",  gt_ei_all, "num2", d_ei_all, "num2",
             goodlow=True, suffix="/1k"),
        _kpi(f"C2 EI Cli-Esc {lw}",  gt_ei_cli, "num2", goodlow=True, suffix="/1k"),
        _kpi(f"Total Esc {lw}",       gt_esc,    "int",  goodlow=True),
        _kpi(f"Zones tracked",
             len([z for z in rows_all if z!="Grand Total"]), "int", goodlow=False),
    ]

    zones = [z for z in rows_all if z != "Grand Total"]
    ei_all_vals = [_latest_val(rows_all,z,weeks_a,ei_k)  for z in zones]
    ei_cli_vals = [_latest_val(rows_cli,z,weeks_c,ei_k)  for z in zones]
    esc_vals    = [_latest_val(rows_all,z,weeks_a,"Total Esc") for z in zones]

    # Simple EI time-series from the C2 EI simple block
    ei_row_labs = [l for l in rows_ei]
    ei_trend_ds = [(l, [rows_ei.get(l,{}).get(w) for w in ei_wks]) for l in ei_row_labs]

    charts = [
        _bar(f"C2 EI by Zone {lw}", "num2", zones,
             [("All-Source EI", ei_all_vals), ("Client-Esc EI", ei_cli_vals)]),
        _line("C2 EI Trend", "num2", ei_wks, ei_trend_ds),
        _bar(f"Total Esc by Zone {lw}", "int", zones, [("Total Esc", esc_vals)]),
    ]

    tbl_rows = [[z,
                 _latest_val(rows_all,z,weeks_a,ei_k),
                 _latest_val(rows_cli,z,weeks_c,ei_k),
                 _latest_val(rows_all,z,weeks_a,"Total Esc")]
                for z in zones + ["Grand Total"] if z in rows_all]
    tables = [_table(f"C2 Zone Summary {lw}",
                     ["Zone","All-Src EI","Cli-Esc EI","Total Esc"],
                     tbl_rows, ["t","num2","num2","int"])]

    return {
        "id":"c2","name":"Zone Wise - C2","icon":"C2",
        "subtitle":"C2 Escalation Index by zone",
        "latest_week": lw,
        "narrative":(f"{lw}: C2 All-Source EI = {gt_ei_all:.2f}/1k."
                     if gt_ei_all else f"{lw}: C2 EI data."),
        "kpis":kpis,"charts":charts,"tables":tables,
    }


# ─── Sheet 9: RTO & Cancelled Reasons ─────────────────────────────────────
def parse_rto(wb):
    ws = wb["RTO & Cancelled Reasons"]
    rows_rto, rto_cols = read_simple(ws, 0, 1, 39, c0=0)
    rows_can, can_cols = read_simple(ws, 0, 1, 8,  c0=6)

    # Grand Total col
    gt_col = "Grand Total" if "Grand Total" in rto_cols else rto_cols[-1]

    rto_data = [(l, rows_rto[l].get(gt_col))
                for l in rows_rto
                if l not in ("Grand Total",) and rows_rto[l].get(gt_col) is not None]
    can_data = [(l, rows_can[l].get(gt_col))
                for l in rows_can
                if l not in ("Grand Total",) and rows_can[l].get(gt_col) is not None]

    rto_data_s = sorted(rto_data, key=lambda x: x[1] or 0, reverse=True)[:10]
    can_data_s = sorted(can_data, key=lambda x: x[1] or 0, reverse=True)[:8]

    gt_rto = rows_rto.get("Grand Total",{}).get(gt_col)
    gt_can = rows_can.get("Grand Total",{}).get(gt_col)

    kpis = [
        _kpi("Top RTO Reason",       rto_data_s[0][1] if rto_data_s else None,
             "int", goodlow=True, text=rto_data_s[0][0] if rto_data_s else ""),
        _kpi("RTO Grand Total",      gt_rto, "int", goodlow=True),
        _kpi("Top Cancelled Reason", can_data_s[0][1] if can_data_s else None,
             "int", goodlow=True, text=can_data_s[0][0] if can_data_s else ""),
        _kpi("Cancelled Grand Total",gt_can, "int", goodlow=True),
    ]

    rto_labs = [x[0] for x in rto_data_s]
    rto_vals = [x[1] for x in rto_data_s]
    can_labs = [x[0] for x in can_data_s]
    can_vals = [x[1] for x in can_data_s]

    charts = [
        _bar("Top 10 RTO Reasons", "int", rto_labs,
             [("Grand Total", rto_vals)], horizontal=True),
        _bar("Top 8 Cancelled Reasons", "int", can_labs,
             [("Grand Total", can_vals)], horizontal=True),
    ]

    wk_cols = [c for c in rto_cols if c != "Grand Total"]
    rto_tbl = [[l] + [rows_rto.get(l,{}).get(c) for c in wk_cols + [gt_col]]
               for l,_ in rto_data_s]
    can_tbl = [[l] + [rows_can.get(l,{}).get(c) for c in wk_cols[:2] + [gt_col]]
               for l,_ in can_data_s]

    tables = [
        _table("RTO Reasons", ["Reason"] + wk_cols + [gt_col],
               rto_tbl, ["t"] + ["int"] * (len(wk_cols)+1)),
        _table("Cancelled Reasons", ["Reason"] + wk_cols[:2] + [gt_col],
               can_tbl, ["t","int","int","int"]),
    ]

    return {
        "id":"rto","name":"RTO & Cancelled Reasons","icon":"RC",
        "subtitle":"Top return & cancellation drivers",
        "latest_week": gt_col,
        "narrative":(f"Grand Total: RTO {int(gt_rto or 0):,} orders; "
                     f"top reason: {rto_data_s[0][0] if rto_data_s else 'N/A'}."),
        "kpis":kpis,"charts":charts,"tables":tables,
    }


# ─── Sheet 10: Action User ─────────────────────────────────────────────────
def parse_action(wb):
    ws = wb["Action User"]
    rows_can, can_cols = read_simple(ws, 1, 2, 100, c0=0)
    rows_rto, rto_cols = read_simple(ws, 1, 2, 100, c0=6)

    rows_can = {k: v for k, v in rows_can.items()
                if any(x is not None for x in v.values())}
    rows_rto = {k: v for k, v in rows_rto.items()
                if any(x is not None for x in v.values())}

    gt_col_c = "Grand Total" if "Grand Total" in can_cols else can_cols[-1] if can_cols else None
    gt_col_r = "Grand Total" if "Grand Total" in rto_cols else rto_cols[-1] if rto_cols else None

    can_data = [(l, rows_can[l].get(gt_col_c))
                for l in rows_can if rows_can[l].get(gt_col_c) is not None]
    rto_data = [(l, rows_rto[l].get(gt_col_r))
                for l in rows_rto if rows_rto[l].get(gt_col_r) is not None]

    can_s = sorted(can_data, key=lambda x: x[1] or 0, reverse=True)[:10]
    rto_s = sorted(rto_data, key=lambda x: x[1] or 0, reverse=True)[:10]

    gt_can = sum((v or 0) for _,v in can_data)
    gt_rto = sum((v or 0) for _,v in rto_data)

    kpis = [
        _kpi("Top Cancel User",   can_s[0][1] if can_s else None,
             "int", goodlow=True, text=can_s[0][0] if can_s else ""),
        _kpi("Cancel Total",      gt_can, "int", goodlow=True),
        _kpi("Top RTO User",      rto_s[0][1] if rto_s else None,
             "int", goodlow=True, text=rto_s[0][0] if rto_s else ""),
        _kpi("RTO Total",         gt_rto, "int", goodlow=True),
    ]

    charts = [
        _bar("Top 10 Cancellation Users", "int",
             [x[0] for x in can_s], [("Grand Total",[x[1] for x in can_s])],
             horizontal=True),
        _bar("Top 10 RTO Action Users", "int",
             [x[0] for x in rto_s], [("Grand Total",[x[1] for x in rto_s])],
             horizontal=True),
    ]

    can_wks = [c for c in can_cols if c != gt_col_c][:2]
    rto_wks = [c for c in rto_cols if c != gt_col_r][:2]
    can_tbl = [[l] + [rows_can.get(l,{}).get(c) for c in can_wks + [gt_col_c]]
               for l,_ in can_s]
    rto_tbl = [[l] + [rows_rto.get(l,{}).get(c) for c in rto_wks + [gt_col_r]]
               for l,_ in rto_s]

    tables = [
        _table("Cancellation Action Users", ["User"] + can_wks + [gt_col_c],
               can_tbl, ["t","int","int","int"]),
        _table("RTO Action Users", ["User"] + rto_wks + [gt_col_r],
               rto_tbl, ["t","int","int","int"]),
    ]

    return {
        "id":"action","name":"Action User","icon":"AU",
        "subtitle":"Top users driving cancellations & RTOs",
        "latest_week": gt_col_c or "N/A",
        "narrative":(f"Top cancellation user: {can_s[0][0] if can_s else 'N/A'} "
                     f"({int(can_s[0][1] or 0):,}). "
                     f"Top RTO user: {rto_s[0][0] if rto_s else 'N/A'}."),
        "kpis":kpis,"charts":charts,"tables":tables,
    }


# ─── Sheet 11: POST UDS FAKE ───────────────────────────────────────────────
def parse_postfake(wb):
    ws = wb["POST UDS FAKE"]
    W = WINDOW["postfake"]
    rows, weeks, subs = read_block(ws, 0, 1, 2, 20, c0=0, maxgroups=W)
    rows = nonempty_rows(rows)
    lw = weeks[-1] if weeks else "?"

    awb_k  = "Awb Count" if "Awb Count" in subs else subs[0] if subs else "Awb Count"
    fake_k = "Fake"      if "Fake" in subs      else subs[1] if len(subs)>1 else "Fake"
    pct_k  = "%"         if "%" in subs          else subs[2] if len(subs)>2 else "%"

    gt_awb  = _latest_val(rows,"Grand Total",weeks,awb_k)
    gt_fake = _latest_val(rows,"Grand Total",weeks,fake_k)
    gt_pct  = _latest_val(rows,"Grand Total",weeks,pct_k)
    d_pct   = _delta(rows,"Grand Total",pct_k,weeks)

    kpis = [
        _kpi(f"Fake% GT {lw}",   gt_pct,  "pct2", d_pct, "pp", goodlow=True),
        _kpi(f"Fake Count {lw}", gt_fake, "int",  goodlow=True),
        _kpi(f"AWB Count {lw}",  gt_awb,  "int",  goodlow=False),
        _kpi("PODs tracked",
             len([l for l in rows if l != "Grand Total"]), "int", goodlow=False),
    ]

    pod_labs  = [l for l in rows if l != "Grand Total"]
    pct_vals  = [_latest_val(rows,l,weeks,pct_k)  for l in pod_labs]
    fake_vals = [_latest_val(rows,l,weeks,fake_k) for l in pod_labs]

    trend_ds = [("Fake%",     _series(rows,"Grand Total",weeks,pct_k)),
                ("Fake Count",_series(rows,"Grand Total",weeks,fake_k))]

    charts = [
        _bar(f"Fake% by POD {lw}", "pct2", pod_labs, [("Fake%", pct_vals)]),
        _line("Fake% Grand Total Trend", "pct2", weeks, [("Fake%", _series(rows,"Grand Total",weeks,pct_k))]),
        _bar(f"Fake Count by POD {lw}", "int", pod_labs, [("Fake", fake_vals)]),
    ]

    tbl_rows = [[l, _latest_val(rows,l,weeks,awb_k),
                 _latest_val(rows,l,weeks,fake_k),
                 _latest_val(rows,l,weeks,pct_k)]
                for l in pod_labs + ["Grand Total"] if l in rows]
    tables = [_table(f"POST UDS Fake {lw}",
                     ["POD","AWB Count","Fake","Fake%"],
                     tbl_rows, ["t","int","int","pct2"])]

    return {
        "id":"postfake","name":"POST UDS FAKE","icon":"PF",
        "subtitle":"Post-UDS fake delivery rates by POD",
        "latest_week": lw,
        "narrative":(f"{lw}: Grand Total Fake% = {(gt_pct or 0)*100:.2f}% "
                     f"({(d_pct or 0)*100:+.2f}pp vs prior)."),
        "kpis":kpis,"charts":charts,"tables":tables,
    }


# ─── Sheet 12: Pre & Post_Data ─────────────────────────────────────────────
def parse_prepost(wb):
    ws = wb["Pre & Post_Data"]
    # Read all rows we need (rows 0-36, 0-indexed)
    all_rows = get_rows(ws, 37)

    def cell(ri, ci):
        r = all_rows[ri] if ri < len(all_rows) else ()
        return num(r[ci]) if ci < len(r) else None

    def label(ri, ci):
        r = all_rows[ri] if ri < len(all_rows) else ()
        v = r[ci] if ci < len(r) else None
        if v is None: return None
        s = str(v).strip()
        return s if s else None

    # Forward POD block: label-col=0, data rows 3-22
    fwd_pods = []
    for ri in range(3, 23):
        lbl = label(ri, 0)
        if not lbl or lbl in SKIP_LABELS: continue
        fwd_pods.append({
            "label": lbl,
            "PreDel%":  cell(ri, 3),
            "PostDel%": cell(ri, 8),
            "PreRet%":  cell(ri, 5),
            "PostRet%": cell(ri, 10),
        })

    # Reverse POD block: label-col=12, data rows 3-22
    rev_pods = []
    for ri in range(3, 23):
        lbl = label(ri, 12)
        if not lbl or lbl in SKIP_LABELS: continue
        rev_pods.append({
            "label": lbl,
            "PrePick%":  cell(ri, 15),
            "PostPick%": cell(ri, 20),
        })

    # Grand Total row
    gt_fwd = next((r for r in fwd_pods if r["label"] == "Grand Total"), None)
    gt_rev = next((r for r in rev_pods if r["label"] == "Grand Total"), None)

    gt_post_del  = gt_fwd["PostDel%"]  if gt_fwd else None
    gt_pre_del   = gt_fwd["PreDel%"]   if gt_fwd else None
    gt_post_pick = gt_rev["PostPick%"] if gt_rev else None
    gt_pre_pick  = gt_rev["PrePick%"]  if gt_rev else None

    kpis = [
        _kpi("Pre Del%  (Fwd GT)",  gt_pre_del,   "pct2", goodlow=False),
        _kpi("Post Del% (Fwd GT)",  gt_post_del,  "pct2",
             (gt_post_del - gt_pre_del) if gt_post_del and gt_pre_del else None,
             "pp", goodlow=False),
        _kpi("Pre Pick% (Rev GT)",  gt_pre_pick,  "pct2", goodlow=False),
        _kpi("Post Pick% (Rev GT)", gt_post_pick, "pct2",
             (gt_post_pick - gt_pre_pick) if gt_post_pick and gt_pre_pick else None,
             "pp", goodlow=False),
    ]

    non_gt_fwd = [r for r in fwd_pods if r["label"] not in ("Grand Total","(blank)")]
    non_gt_rev = [r for r in rev_pods if r["label"] not in ("Grand Total","(blank)")]

    fwd_labs = [r["label"] for r in non_gt_fwd]
    rev_labs = [r["label"] for r in non_gt_rev]

    charts = [
        _bar("Fwd – Pre vs Post Delivery%", "pct2", fwd_labs,
             [("Pre Del%",  [r["PreDel%"]  for r in non_gt_fwd]),
              ("Post Del%", [r["PostDel%"] for r in non_gt_fwd])]),
        _bar("Rev – Pre vs Post Pick%", "pct2", rev_labs,
             [("Pre Pick%",  [r["PrePick%"]  for r in non_gt_rev]),
              ("Post Pick%", [r["PostPick%"] for r in non_gt_rev])]),
    ]

    tbl_rows_f = [[r["label"], r["PreDel%"], r["PostDel%"],
                   (r["PostDel%"]-r["PreDel%"])
                   if r["PostDel%"] is not None and r["PreDel%"] is not None else None]
                  for r in fwd_pods]
    tbl_rows_r = [[r["label"], r["PrePick%"], r["PostPick%"],
                   (r["PostPick%"]-r["PrePick%"])
                   if r["PostPick%"] is not None and r["PrePick%"] is not None else None]
                  for r in rev_pods]

    tables = [
        _table("Fwd Pre vs Post", ["POD","Pre Del%","Post Del%","Δpp"],
               tbl_rows_f, ["t","pct2","pct2","pp"]),
        _table("Rev Pre vs Post", ["POD","Pre Pick%","Post Pick%","Δpp"],
               tbl_rows_r, ["t","pct2","pct2","pp"]),
    ]

    uplift = ((gt_post_del or 0) - (gt_pre_del or 0))
    return {
        "id":"prepost","name":"Pre & Post_Data","icon":"PP",
        "subtitle":"Pre-UDS vs Post-UDS delivery improvement",
        "latest_week": "Wk 11&12 vs 8&9",
        "narrative":(f"Post-UDS Del% = {(gt_post_del or 0)*100:.1f}% "
                     f"vs Pre-UDS {(gt_pre_del or 0)*100:.1f}% "
                     f"({uplift*100:+.1f}pp uplift)."),
        "kpis":kpis,"charts":charts,"tables":tables,
    }


# ─── Auto-detect source ────────────────────────────────────────────────────
def find_source():
    here   = os.path.dirname(os.path.abspath(__file__))
    parent = os.path.dirname(here)
    candidates = []
    for d in (here, parent):
        for f in glob.glob(os.path.join(d, "UDS_PPT_Data_Wk_*.xlsx")):
            bn = os.path.basename(f)
            if "Dashboard" not in bn and not bn.startswith("~$"):
                candidates.append(f)
    if not candidates:
        raise FileNotFoundError(
            "No UDS_PPT_Data_Wk_*.xlsx found. "
            "Place it in the same folder as this script or its parent."
        )
    return max(candidates, key=os.path.getmtime)


# ─── Main ──────────────────────────────────────────────────────────────────
def main():
    src = find_source()
    print(f"Source: {src}")
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)

    parsers = [
        parse_cci, parse_ei, parse_source, parse_client,
        parse_pod, parse_issue, parse_zone, parse_c2,
        parse_rto, parse_action, parse_postfake, parse_prepost,
    ]

    sheets = []
    for fn in parsers:
        try:
            s = fn(wb)
            sheets.append(s)
            print(f"  OK {s['name']:40s}  latest={s.get('latest_week','?')}")
        except Exception as e:
            import traceback
            print(f"  FAIL {fn.__name__}: {e}")
            traceback.print_exc()

    wb.close()

    dashboard = {
        "meta": {
            "title":     "UDS Weekly Operations Dashboard",
            "subtitle":  "Logistics Performance Analytics",
            "source":    os.path.basename(src),
            "generated": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "nsheets":   len(sheets),
        },
        "sheets": sheets,
    }

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data.js")
    js  = "window.DASHBOARD = " + json.dumps(clean_nan(dashboard),
                                               ensure_ascii=False, indent=2) + ";"
    with open(out, "w", encoding="utf-8") as f:
        f.write(js)

    total_charts = sum(len(s.get("charts", [])) for s in sheets)
    total_kpis   = sum(len(s.get("kpis",   [])) for s in sheets)
    print(f"\nOK data.js written -> {out}")
    print(f"  {len(sheets)} sheets | {total_kpis} KPIs | {total_charts} charts")


if __name__ == "__main__":
    main()
